from __future__ import annotations

import base64
import os
import re
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document
from dotenv import load_dotenv


@dataclass
class LoadedDocument:
    file_name: str
    file_type: str
    text: str
    pages: list[dict]


PDF_MIN_TEXT_CHARS = 50
OCR_MODEL = "PaddlePaddle/PaddleOCR-VL-1.5"


def load_document(path: str | Path) -> LoadedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)
    if suffix in {".txt", ".md"}:
        return _load_text(file_path)
    raise ValueError(f"Unsupported file type: {suffix}")


def _load_pdf(path: Path) -> LoadedDocument:
    ocr_doc = _load_pdf_by_ocr(path)
    if len(ocr_doc.text.strip()) >= PDF_MIN_TEXT_CHARS:
        return ocr_doc

    text_doc = _load_pdf_by_pymupdf(path)
    if len(text_doc.text.strip()) >= PDF_MIN_TEXT_CHARS:
        return text_doc

    return LoadedDocument(
        path.name,
        "pdf",
        "PDF 文本提取失败：OCR 与 PyMuPDF 文本层均未提取到足够内容。可能是低清扫描件、加密文件或图片质量过低。",
        [],
    )


def _load_pdf_by_pymupdf(path: Path) -> LoadedDocument:
    pages = []
    with fitz.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append({"page": index, "text": text})
    full_text = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
    return LoadedDocument(path.name, "pdf", full_text, pages)


def _load_pdf_by_ocr(path: Path) -> LoadedDocument:
    load_dotenv(dotenv_path=Path(".env"))
    api_key = os.getenv("OPENAI_API_KEY")
    base_url = os.getenv("OPENAI_BASE_URL")
    model = os.getenv("OCR_MODEL", OCR_MODEL)
    if not api_key or not base_url:
        return LoadedDocument(path.name, "pdf", "", [])

    try:
        from openai import OpenAI

        client = OpenAI(
            api_key=api_key,
            base_url=base_url,
            timeout=float(os.getenv("OCR_TIMEOUT", os.getenv("OPENAI_TIMEOUT", "180"))),
            max_retries=int(os.getenv("OPENAI_MAX_RETRIES", "0")),
        )
        pages = []
        with fitz.open(path) as doc:
            for index, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_b64 = base64.b64encode(pix.tobytes("png")).decode("ascii")
                text = _ocr_page(client, model, image_b64)
                text = _clean_ocr_text(text)
                if text:
                    pages.append({"page": index, "text": text, "source": "ocr"})
        full_text = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
        return LoadedDocument(path.name, "pdf", full_text, pages)
    except Exception:
        return LoadedDocument(path.name, "pdf", "", [])


def _ocr_page(client, model: str, image_b64: str) -> str:
    prompt = (
        "你是 OCR 引擎。请识别这张中文简历页面里的可见文字。"
        "只输出纯文本，不要解释，不要总结，不要 Markdown。"
        "不要输出坐标、标签、<|LOC|>、公式、无关外文乱码。"
        "如果某一行无法识别，直接跳过，不要猜测。"
        "尽量保持原始阅读顺序和换行。"
    )
    response = client.chat.completions.create(
        model=model,
        temperature=0,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{image_b64}"},
                    },
                ],
            }
        ],
    )
    return response.choices[0].message.content or ""


def _clean_ocr_text(text: str) -> str:
    text = re.sub(r"<\|LOC_\d+\|>", "", text)
    text = re.sub(r"\\\((.*?)\\\)", r"\1", text)
    text = text.replace("\\_", "_")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _load_docx(path: Path) -> LoadedDocument:
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return LoadedDocument(path.name, "docx", "\n".join(parts), [])


def _load_text(path: Path) -> LoadedDocument:
    return LoadedDocument(path.name, path.suffix.lower().strip("."), path.read_text(encoding="utf-8", errors="ignore"), [])
